# =============================================================================
# Drafts socioeconomic_poverty_area_map.csv from the PSA poverty workbook and the
# PSGC datafile.
#
# What it does
#   Lists the areas published in the poverty workbook (sheet tab2), gives each one a
#   level, a parent region (as published), a PSGC code, a stable area_key and a
#   proposed geographic coverage basis.
#
# This is a DRAFTING tool, not a pipeline step
#   The CSV it writes is a reviewed contract kept in the repo. Run this script when PSA
#   publishes a new release, compare the output with the CSV in the repo, and review
#   every difference before committing. The silver job reads the CSV, never this script.
#
# How to run (laptop, or a Databricks notebook cell with the Volume paths)
#   python build_socioeconomic_poverty_area_map.py \
#       --poverty  <landing_root>/socioeconomic/psa_poverty/psa_poverty_province_biennial_2021_2025.xlsx \
#       --psgc     <landing_root>/geography/psa_psgc/psa_psgc_barangay_asof_20260713.xlsx \
#       --out      socioeconomic_poverty_area_map.csv
#   Then copy the reviewed CSV to <landing_root>/reference/socioeconomic/, which is the folder
#   socioeconomic_silver_00_reference.sql reads.
#
# What checks the codes afterwards
#   This script takes PSGC codes from the PSGC datafile. The codes are then validated against
#   public_health.silver.geography_silver by socioeconomic_silver_05_poverty_geography_check.sql
#   and by socioeconomic_silver_qc.py (S04, S93).
#
# 2026-10-11: Isabela City now maps to 0990101000 (the PSGC city row), the same code the diseases
#   area map uses. Every code written must now sit at the expected PSGC level.
#
# Requires: openpyxl.
# =============================================================================

import argparse
import collections
import csv
import re

import openpyxl

DATA_START_ROW = 7  # 0-based, same as data_start_row in sources_socioeconomic.yaml
YEARS = (2021, 2023, 2025)
INCIDENCE_COL = {2021: 4, 2023: 5, 2025: 6}  # positions in tab2 (poverty_tab2 layout in the YAML)

# Manual decisions. Every entry here is a judgement that a person reviewed.
PROVINCE_NAME_FIXES = {"mt. province": "mountain province"}  # poverty name -> PSGC name
LISTED_CITIES = {
    "Isabela City": "0990101000",  # PSGC city row "City of Isabela". Not 0990100000
    # "City of Isabela (Not a Province)", a placeholder
    # with no level and no population
    "Cotabato City": "1908703000",
}  # PSGC "City of Cotabato"
SUPERSEDED = {"Maguindanao": "hist:maguindanao"}  # split into del Norte / del Sur; no current code

FIELDS = [
    "area_key",
    "geo_name_source",
    "geo_level",
    "psgc_code",
    "parent_geo_name_source",
    "geo_coverage_basis",
    "coverage_confirmed",
    "first_estimate_year",
    "is_superseded_area",
    "comparison_eligible",
    "exclusion_reason",
    "source_order",
]


def norm(text):
    # Lower case, single spaces: used to compare names.
    return re.sub(r"\s+", " ", str(text)).strip().lower()


def read_poverty_areas(path):
    # Area rows of tab2, selected by STRUCTURE: every row from the first data row up to the
    # "Notes:" line that has an area name. No row is chosen by whether a value is a number.
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["tab2"].iter_rows(values_only=True))
    wb.close()
    areas = []
    for i in range(DATA_START_ROW, len(rows)):
        name = rows[i][0]
        if isinstance(name, str) and name.strip().lower().startswith("notes"):
            break
        if isinstance(name, str) and name.strip():
            areas.append((name.strip(), rows[i]))
    return areas


def read_psgc(path):
    # Region and province names with their codes, plus the population note of each province
    # (the unnamed column after "2024 Population", e.g. "(excluding CITY OF BAGUIO)").
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["PSGC"].iter_rows(values_only=True))[1:]
    wb.close()
    regions = {norm(r[1]): str(r[0]) for r in rows if r[3] == "Reg"}
    provinces = {norm(r[1]): (str(r[0]), r[9]) for r in rows if r[3] == "Prov"}
    levels = {str(r[0]): r[3] for r in rows}  # code -> PSGC level (Reg, Prov, City, ...)
    return regions, provinces, levels


def build(areas, regions, provinces, levels):
    out, parent = [], ""
    for order, (name, row) in enumerate(areas, 1):
        n = norm(name)
        code, basis, eligible, reason, superseded = "", "", "true", "", "false"

        if name == "PHILIPPINES":
            level, key, basis = "country", "ph", "country"
        elif n in regions:
            level, code, basis = "region", regions[n], "region"
            key, parent = "psgc:" + code, name  # areas below belong to this region
        elif re.match(r"^\d(st|nd|rd|th) District$", name):
            level, key, basis = "ncr_district", "ncr-district:" + name[0], "ncr_district"
            eligible, reason = "false", "no_psgc_equivalent"
        elif name in LISTED_CITIES:
            level, code, basis = "city", LISTED_CITIES[name], "city_listed_separately"
            key = "psgc:" + code
            if levels.get(code) != "City":  # a typed code must be a PSGC city row
                raise ValueError(
                    f"{name}: code {code} is not a City row in the PSGC file "
                    f"(found level: {levels.get(code)!r})"
                )
        elif name in SUPERSEDED:
            level, key, basis = "province", SUPERSEDED[name], "superseded_province"
            eligible, reason, superseded = "false", "superseded_area", "true"
        else:
            level = "province"
            match = provinces.get(PROVINCE_NAME_FIXES.get(n, n))
            if match is None:  # a new or renamed area: stop, do not guess
                raise ValueError(
                    f"'{name}' is not a PSGC region or province and has no manual rule. "
                    "Add it to PROVINCE_NAME_FIXES, LISTED_CITIES or SUPERSEDED."
                )
            code, note = match
            key = "psgc:" + code
            if note and "ISABELA" in note.upper():
                basis, eligible, reason = (
                    "province_excl_listed_city",
                    "false",
                    "coverage_unconfirmed",
                )
            elif note:  # PSGC population excludes a highly urbanized city
                basis, eligible = "province_incl_huc", "false"
                reason = "poverty_includes_huc_psgc_population_excludes"
            else:
                basis = "province"

        # First year with a published incidence estimate.
        first_year = next(
            (y for y in YEARS if isinstance(row[INCIDENCE_COL[y]], (int, float))), None
        )
        if first_year is None:
            raise ValueError(f"'{name}' has no incidence estimate in any year")

        out.append(
            {
                "area_key": key,
                "geo_name_source": name,
                "geo_level": level,
                "psgc_code": code,
                "parent_geo_name_source": "" if level in ("country", "region") else parent,
                "geo_coverage_basis": basis,
                "coverage_confirmed": "false" if level == "province" else "true",
                "first_estimate_year": first_year,
                "is_superseded_area": superseded,
                "comparison_eligible": eligible,
                "exclusion_reason": reason,
                "source_order": order,
            }
        )
    return out


def check(rows):
    # Stop on duplicate keys or codes; print a summary to review.
    for field in ("area_key", "geo_name_source"):
        dup = [k for k, v in collections.Counter(r[field] for r in rows).items() if v > 1]
        if dup:
            raise ValueError(f"duplicate {field}: {dup}")
    dup = [
        k
        for k, v in collections.Counter(r["psgc_code"] for r in rows if r["psgc_code"]).items()
        if v > 1
    ]
    if dup:
        raise ValueError(f"duplicate psgc_code: {dup}")
    area_years = sum(len([y for y in YEARS if y >= r["first_estimate_year"]]) for r in rows)
    print(f"areas: {len(rows)}  expected area-years: {area_years}")
    print("by level:   ", dict(collections.Counter(r["geo_level"] for r in rows)))
    print("by coverage:", dict(collections.Counter(r["geo_coverage_basis"] for r in rows)))
    print(
        "not comparison-eligible:",
        dict(
            collections.Counter(
                r["exclusion_reason"] for r in rows if r["comparison_eligible"] == "false"
            )
        ),
    )


def main():
    p = argparse.ArgumentParser(
        description="Draft the poverty area map from the PSA poverty and PSGC files"
    )
    p.add_argument("--poverty", required=True, help="Path to the PSA poverty workbook")
    p.add_argument("--psgc", required=True, help="Path to the PSGC datafile")
    p.add_argument("--out", default="socioeconomic_poverty_area_map.csv", help="CSV to write")
    args = p.parse_args()

    areas = read_poverty_areas(args.poverty)
    regions, provinces, levels = read_psgc(args.psgc)
    rows = build(areas, regions, provinces, levels)
    check(rows)
    with open(args.out, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    print(f"wrote {args.out}")


if __name__ == "__main__":
    main()
